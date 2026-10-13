from __future__ import annotations

import json
import shutil
import stat
from io import BytesIO
from decimal import Decimal
from pathlib import Path
import openpyxl
from .store import dump, fingerprint, now


XFT_FIELDS = {1:"bill_no", 3:"title", 4:"tax_no", 5:"address", 6:"phone", 7:"bank", 8:"bank_account",
              9:"natural", 10:"seller_name", 11:"seller_address", 12:"seller_phone", 13:"seller_bank", 14:"seller_account",15:"invoice_type", 18:"remark", 23:"email", 24:"tax_code",
              25:"invoice_name", 26:"unit", 27:"spec", 28:"qty", 29:"price", 30:"rate_text", 31:"policy",
              32:"amount", 33:"included", 36:"name"}

ALIASES = {
 "bill_no": ["序号", "发票流水号", "开票流水号", "单据号", "发票序号", "关联序号"],
 "title": ["购买方名称", "购方名称", "购买方纳税人名称"],
 "tax_no": ["购买方识别号（统一社会信用代码）", "购买方纳税人识别号", "购买方纳税人识别号（统一社会信用代码）", "购买方识别号", "购方税号"],
 "invoice_type": ["发票类型", "发票种类"], "address": ["购买方地址"], "phone": ["购买方电话", "购买方手机号"],
 "email": ["购买方邮箱", "购买方电子邮箱"], "bank": ["购买方开户银行"], "bank_account": ["购买方银行账号"],
 "remark": ["备注"], "gross": ["价税合计"], "invoice_name": ["项目名称", "货物或应税劳务、服务名称", "商品名称"],
 "tax_code": ["商品和服务税收分类编码", "税收分类编码"], "spec": ["规格型号"], "unit": ["单位", "计量单位"],
 "qty": ["数量"], "net_price": ["单价（不含税）", "单价"], "net": ["金额（不含税）", "金额"],
 "rate": ["税率", "税率/征收率"], "tax": ["税额"], "included_flag": ["含税标志"],
}


def clean_header(value):
    return str(value or "").replace("*", "").replace(" ", "").replace("(", "（").replace(")", "）").strip()


def signature(path, profile):
    wb = openpyxl.load_workbook(BytesIO(Path(path).read_bytes()), read_only=False, data_only=False)
    try:
        info = {"sheets": wb.sheetnames, "tables": []}
        for table in profile["tables"]:
            ws = wb[table["sheet"]]
            info["tables"].append({"sheet": ws.title, "row": table["header_row"],
                "headers": [clean_header(c.value) for c in ws[table["header_row"]]],
                "validation": [(str(v.sqref), v.type, v.formula1, v.formula2) for v in ws.data_validations.dataValidation]})
        return fingerprint(info)
    finally:
        wb.close()


def inspect_template(target, path):
    wb = openpyxl.load_workbook(BytesIO(Path(path).read_bytes()), read_only=False, data_only=False)
    try:
        if target == "xinfutong":
            if "导入模板" not in wb.sheetnames:
                raise ValueError("新福通模板缺少「导入模板」页签")
            heads = [str(c.value or "") for c in wb["导入模板"][2]]
            checks = {1:"单据号", 3:"购买方", 10:"法人公司", 24:"税收", 28:"数量", 32:"金额"}
            if len(heads) < 36 or any(v not in heads[k-1] for k, v in checks.items()):
                raise ValueError("新福通模板表头已变化，需要更新适配器；未套用旧列号")
            return {"tables": [{"sheet":"导入模板", "header_row":2, "kind":"combined",
                                 "columns":{str(k):v for k,v in XFT_FIELDS.items()}}]}
        if "excelVersion" in wb.sheetnames and wb["excelVersion"]["A4"].value == "pt:20260401":
            expected = {"1-发票基本信息":{1:"发票流水号",2:"发票类型",4:"是否含税",5:"受票方自然人标识",6:"购买方名称",7:"购买方纳税人识别号",31:"购买方邮箱"},
                "2-发票明细信息":{1:"发票流水号",2:"项目名称",3:"商品和服务税收编码",6:"数量",7:"单价",8:"金额",9:"税率"}}
            for name,columns in expected.items():
                if name not in wb.sheetnames or any(clean_header(wb[name].cell(3,col).value)!=clean_header(value) for col,value in columns.items()):
                    raise ValueError("V260401官方模板结构与已适配版本不符，请更新适配器")
            return {"version":"pt:20260401","text_input":True,"max_file_lines":2000,"max_file_invoices":998,
                "tables":[
                    {"sheet":"1-发票基本信息","header_row":3,"kind":"base","columns":{str(k):v for k,v in {
                        1:"bill_no",2:"invoice_type",4:"included",5:"natural",6:"title",7:"tax_no",11:"address",17:"phone",18:"bank",19:"bank_account",23:"remark",28:"seller_bank",29:"seller_account",31:"email"}.items()}},
                    {"sheet":"2-发票明细信息","header_row":3,"kind":"detail","columns":{str(k):v for k,v in {
                        1:"bill_no",2:"invoice_name",3:"tax_code",4:"spec",5:"unit",6:"qty",7:"price",8:"amount",9:"rate",11:"preferential"}.items()}},
                    {"sheet":"3-特定业务信息","header_row":3,"kind":"unused","columns":{}},
                    {"sheet":"4-附加要素信息","header_row":3,"kind":"unused","columns":{}}]}
        tables = []
        aliases = {clean_header(a):key for key, values in ALIASES.items() for a in values}
        for ws in wb:
            best = None
            for row in ws.iter_rows(max_row=min(ws.max_row, 15)):
                cols = {str(c.column):aliases[clean_header(c.value)] for c in row if clean_header(c.value) in aliases}
                fields = set(cols.values())
                kind = "detail" if {"bill_no","invoice_name","tax_code","qty","net","rate"} <= fields else "base" if {"bill_no","title","tax_no"} <= fields else None
                if kind and (best is None or len(cols) > len(best["columns"])):
                    unknown = [str(c.value) for c in row if c.value and "*" in str(c.value) and str(c.column) not in cols]
                    if unknown:
                        raise ValueError("模板存在尚未适配的必填列："+"、".join(unknown))
                    best = {"sheet":ws.title, "header_row":row[0].row, "kind":kind, "columns":cols}
            if best:
                tables.append(best)
        if {t["kind"] for t in tables} != {"base", "detail"}:
            raise ValueError("无法可靠识别浙江税务局模板的基本信息、明细信息和关联编号列；请使用实际下载的官方模板")
        return {"tables":tables, "invoice_types":{"普通":"电子发票（普通发票）","专用":"电子发票（增值税专用发票）"}}
    finally:
        wb.close()


def register(store, target, path):
    if target not in ("xinfutong", "zhejiang"):
        raise ValueError("未知模板类型")
    profile = inspect_template(target, path)
    dest = store.root / "模板" / "v2" / (target + "_" + fingerprint(Path(path).read_bytes().hex())[:12] + ".xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    if Path(path).resolve() != dest.resolve():
        shutil.copyfile(path, dest)
    dest.chmod(stat.S_IREAD | stat.S_IWRITE)
    sig = signature(dest, profile)
    store.execute("INSERT INTO templates VALUES(?,?,?,?,?,?) ON CONFLICT(target) DO UPDATE SET path=excluded.path,signature=excluded.signature,profile=excluded.profile,verified=0,updated=excluded.updated",
                  (target, str(dest), sig, dump(profile), 0, now()))
    store.audit("template", target, None, {"path":str(dest), "signature":sig, "import_verified":False})
    return sig


def render(store, target, invoices, outpath, seller):
    rows = store.rows("SELECT * FROM templates WHERE target=?", (target,))
    if not rows:
        raise ValueError("请在资料设置导入浙江税务局官方模板" if target == "zhejiang" else "请导入新福通官方模板")
    meta = rows[0]
    profile = json.loads(meta["profile"])
    if len(invoices)>profile.get("max_file_invoices",100000) or sum(len(i["payload"]["items"]) for i in invoices)>profile.get("max_file_lines",100000):
        raise ValueError("本批次超过模板单文件安全容量，请按申请分批输出，未生成超限文件")
    if signature(meta["path"], profile) != meta["signature"]:
        raise ValueError("模板结构发生变化，请重新注册模板，未生成该格式")
    wb = openpyxl.load_workbook(BytesIO(Path(meta["path"]).read_bytes()))
    try:
        for table in profile["tables"]:
            ws = wb[table["sheet"]]
            start = table["header_row"] + 1
            for cells in ws.iter_rows(min_row=start):
                for c in cells:
                    if c.__class__.__name__ != "MergedCell":
                        c.value = None
            rownum = start
            if table["kind"] == "unused":continue
            for inv in invoices:
                p = inv["payload"]
                invoice_type = "电子发票（增值税专用发票）" if "专用" in p.get("invoice_type", "") else "电子发票（普通发票）"
                if profile.get("version") == "pt:20260401":
                    invoice_type = "增值税专用发票" if "专" in p.get("invoice_type", "") else "普通发票"
                base = {**p, "bill_no":inv["id"], "natural":"是" if p.get("personal") else "否",
                        "seller_name":seller.get("name", ""), "invoice_type":invoice_type,
                        "seller_address":seller.get("address",""),"seller_phone":seller.get("phone",""),
                        "seller_bank":seller.get("bank",""),"seller_account":seller.get("bank_account",""),
                        "remark":"订单号：" + "，".join(p["order_nos"]), "gross":Decimal(p["amount"])}
                items = [{}] if table["kind"] == "base" else p["items"]
                for item in items:
                    vals = {**base, **item, "policy":"不使用优惠政策", "included":"是", "included_flag":0,"preferential":"否"}
                    if item:
                        for key in ("qty", "amount", "price", "net", "tax", "rate"):
                            vals[key] = Decimal(str(item.get(key, "0")))
                        vals["net_price"] = (vals["net"] / vals["qty"]).quantize(Decimal("0.00000001"))
                        vals["rate_text"] = format((vals['rate']*100).normalize(),"f")+"%"
                    for col, key in table["columns"].items():
                        value=vals.get(key, "")
                        if profile.get("text_input"):
                            value = format(value,"f") if isinstance(value,Decimal) else str(value)
                        cell = ws.cell(rownum, int(col),value)
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                        for dv in ws.data_validations.dataValidation:
                            if cell.coordinate in dv.sqref and dv.type == "list" and dv.formula1 and dv.formula1.startswith('"'):
                                choices = dv.formula1.strip('"').split(',')
                                if cell.value not in (None, "") and str(cell.value) not in choices:
                                    raise ValueError(f"模板数据校验不接受 {key} 的值，请核对业务设置")
                    rownum += 1
        Path(outpath).parent.mkdir(parents=True, exist_ok=True)
        wb.save(outpath)
    finally:
        wb.close()
    # 再读一次，确认产物可打开。
    check = openpyxl.load_workbook(outpath, read_only=True)
    check.close()
    return meta["signature"]
