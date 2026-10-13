import copy
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
import openpyxl

from invoice_assistant.store import Store
from invoice_assistant.domain import validate, identity, split_invoices
from invoice_assistant.engine import Engine, parse_recovery_files
from invoice_assistant.migration import migrate
from invoice_assistant.templates import register

ROOT=Path(__file__).resolve().parents[1]


def request(**overrides):
    raw={"account_id":"account","platform":"pinduoduo","company":"测试公司","shop_id":"123","shop":"测试店",
        "source_id":"app1","applied_at":"2026-10-01 10:00","order_nos":["order1"],"title":"测试买家",
        "tax_no":"91330000TEST000001","personal":False,"invoice_type":"普通","amount":"101.00",
        "status":"pending","tax_included":True,"collection_complete":True,
        "items":[{"order_no":"order1","sku":"sku1","name":"商品A","spec":"大号","qty":"2","amount":"101.00","tax_code":"1090511030000000000","rate":"0.01"}]}
    raw.update(overrides)
    return raw


class V2Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix="invoice-v2-test-")
        self.store=Store(self.temp.name)
        self.store.set("seller",{"name":"测试销售方"})
        self.store.set("default_rate","0.01")
        self.store.set("seller_confirmed",True)
        self.engine=Engine(self.store)
        register(self.store,"xinfutong",ROOT/"模板"/"新福通通用开票申请单模板.xlsx")
        self.bid=self.store.batch("test",{"targets":["xinfutong"]})

    def tearDown(self):self.temp.cleanup()

    def test_decimal_and_idempotent_validation(self):
        p,issues=validate(request(),{},self.engine.settings())
        self.assertFalse(issues)
        self.assertEqual(p["items"][0]["net"],"100.00")
        p2,issues=validate(p,{},self.engine.settings())
        self.assertEqual(p,p2)

    def test_exclusive_amount_does_not_compound(self):
        r=request(tax_included=False)
        r["items"][0]["amount"]="100.00"
        p,issues=validate(r,{},self.engine.settings())
        self.assertFalse(issues)
        self.assertEqual(p["items"][0]["amount"],"101.00")
        p2,issues=validate(p,{},self.engine.settings())
        self.assertEqual(p,p2)

    def test_missing_identity_is_blocked(self):
        p,issues=validate(request(source_id="",applied_at=""),{},self.engine.settings())
        self.assertIn("identity",[i["code"] for i in issues])

    def test_conflict_and_incomplete_block_whole_request(self):
        self.engine.ingest(self.bid,[request(conflict=True),request(source_id="app2",collection_complete=False)])
        s=self.engine.prepare(self.bid,["xinfutong"])
        self.assertEqual(s["files"],[])
        self.assertEqual(s["problem_count"],2)

    def test_same_application_is_not_duplicated(self):
        self.engine.ingest(self.bid,[request()])
        self.engine.ingest(self.bid,[request()])
        self.assertEqual(len(self.engine.requests(self.bid)),1)
        self.engine.prepare(self.bid,["xinfutong"])
        self.engine.prepare(self.bid,["xinfutong"])
        self.assertEqual(len(self.store.rows("SELECT * FROM invoices")),1)
        self.assertEqual(len(self.store.rows("SELECT * FROM exports")),1)

    def test_new_batch_reuses_invoice_identifier(self):
        self.engine.ingest(self.bid,[request()]);self.engine.prepare(self.bid,["xinfutong"])
        bid=self.store.batch("test",{})
        self.engine.ingest(bid,[request()]);self.engine.prepare(bid,["xinfutong"])
        self.assertEqual(len(self.store.rows("SELECT * FROM invoices")),1)

    def test_correction_survives_same_source_refresh(self):
        r=request(title="")
        self.engine.ingest(self.bid,[r])
        self.engine.correct(identity(r),{"title":"已确认买家"})
        self.engine.ingest(self.bid,[r])
        self.assertEqual(self.engine.requests(self.bid)[0]["payload"]["title"],"已确认买家")

    def test_sku_mapping_centralized(self):
        r=request();r["items"][0]["tax_code"]=""
        r2=copy.deepcopy(r);r2["source_id"]="app2"
        self.engine.ingest(self.bid,[r,r2])
        self.engine.correct(identity(r),{"tax_code":"1090511030000000000"},0)
        self.assertTrue(all(not x["problems"] for x in self.engine.requests(self.bid)))
        self.assertEqual(len(self.store.rows("SELECT * FROM mappings")),1)

    def test_amount_not_editable(self):
        r=request();self.engine.ingest(self.bid,[r])
        with self.assertRaises(ValueError):self.engine.correct(identity(r),{"amount":"1"})

    def test_multiple_orders_one_request_one_invoice(self):
        r=request(order_nos=["order1","order2"])
        p,_=validate(r,{},self.engine.settings())
        self.assertEqual(len(split_invoices(identity(r),p,1)),1)

    def test_identical_item_lines_preserved(self):
        r=request(amount="202.00");r["items"]*=2
        p,issues=validate(r,{},self.engine.settings())
        self.assertEqual(len(p["items"]),2)
        self.assertFalse(issues)

    def test_split_partial_result_and_remaining_export(self):
        r=request(amount="202.00");r["items"]*=2
        self.store.set("max_lines",1)
        self.engine.ingest(self.bid,[r]);s=self.engine.prepare(self.bid,["xinfutong"])
        self.assertEqual(s["invoice_count"],2)
        invoices=self.store.rows("SELECT * FROM invoices ORDER BY part")
        self.engine.register_results([{"单据号":invoices[0]["id"],"发票号码":"0000123456789"}])
        self.assertEqual(self.engine.requests(self.bid)[0]["state"],"partially_issued")
        s=self.engine.prepare(self.bid,["xinfutong"])
        self.assertEqual(s["invoice_count"],1)
        self.engine.ingest(self.bid,[r])
        self.assertEqual(self.engine.requests(self.bid)[0]["state"],"partially_issued")
        self.engine.register_results([{"单据号":invoices[1]["id"],"发票号码":"0000123456790"}])
        self.assertEqual(self.engine.requests(self.bid)[0]["state"],"issued")
        self.assertEqual(self.engine.prepare(self.bid,["xinfutong"])["files"],[])

    def test_result_validation_is_atomic(self):
        self.engine.ingest(self.bid,[request()]);self.engine.prepare(self.bid,["xinfutong"])
        inv=self.store.rows("SELECT id FROM invoices")[0]["id"]
        with self.assertRaises(ValueError):self.engine.register_results([{"单据号":inv,"发票号码":"1"},{"单据号":"bad","发票号码":"2"}])
        self.assertEqual(self.store.rows("SELECT state FROM invoices")[0]["state"],"prepared")

    def test_missing_official_zhejiang_is_explicit(self):
        self.engine.ingest(self.bid,[request()])
        s=self.engine.prepare(self.bid,["xinfutong","zhejiang"])
        self.assertEqual(len(s["files"]),1)
        self.assertEqual(s["export_errors"][0]["target"],"zhejiang")

    def test_template_fingerprint_blocks_changed_header(self):
        self.engine.ingest(self.bid,[request()])
        template=self.store.rows("SELECT path FROM templates")[0]["path"]
        wb=openpyxl.load_workbook(template);wb["导入模板"]["A2"]="坏表头";wb.save(template);wb.close()
        s=self.engine.prepare(self.bid,["xinfutong"])
        self.assertFalse(s["files"])
        self.assertIn("结构",s["export_errors"][0]["message"])

    def test_xft_real_template_roundtrip(self):
        self.engine.ingest(self.bid,[request(title="=NOT_FORMULA")])
        s=self.engine.prepare(self.bid,["xinfutong"])
        self.assertEqual(s["export_errors"],[])
        wb=openpyxl.load_workbook(s["files"][0]["path"])
        ws=wb["导入模板"]
        self.assertEqual(ws["C3"].value,"=NOT_FORMULA");self.assertEqual(ws["C3"].data_type,"s")
        self.assertEqual(ws["AF3"].value,101)
        self.assertEqual(ws["X3"].value,"1090511030000000000")
        wb.close()

    def test_migration_idempotent_and_no_samples(self):
        one=migrate(self.store,ROOT);two=migrate(self.store,ROOT)
        self.assertEqual(one,two)
        self.assertFalse(self.store.rows("SELECT * FROM requests"))
        self.assertEqual(len(self.store.rows("SELECT * FROM migrations")),1)

    def test_backup_is_readable(self):
        self.engine.ingest(self.bid,[request()])
        backup=self.store.backup()
        import sqlite3
        from contextlib import closing
        with closing(sqlite3.connect(backup)) as db:self.assertEqual(db.execute("SELECT COUNT(*) FROM requests").fetchone()[0],1)

    def test_known_issued_change_is_blocked(self):
        r=request();self.engine.ingest(self.bid,[r]);self.engine.prepare(self.bid,["xinfutong"])
        inv=self.store.rows("SELECT id FROM invoices")[0]["id"]
        self.engine.register_results([{"单据号":inv,"发票号码":"123"}])
        r["title"]="源数据变更"
        self.engine.ingest(self.bid,[r])
        self.assertIn("issued_change",[p["code"] for p in self.engine.requests(self.bid)[0]["problems"]])


if __name__ == "__main__":unittest.main()
