"""在临时目录验证开票闭环，不写入真实对照库。"""
import copy
import json
import shutil
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from src.config import load_config
from src.ledger import build_ledger, group_invoices, ensure_libs, load_issued, load_customers
from src.pipeline import run, mark_issued, import_missing
from src.readers import parse_file
from src.workflow import assess

BASE = Path(__file__).resolve().parents[1]


def table(path, headers, rows):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def sheet(path, name=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[name] if name else wb.active
        return list(ws.values)
    finally:
        wb.close()


class SilentReporter:
    def __getattr__(self, name):
        return lambda *args, **kwargs: None


class WorkflowTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="invoice_test_")
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.cfg = load_config(BASE, BASE / "tests" / "fixture-config.yaml")
        self.cfg["_base"] = str(self.base)
        for key, path in self.cfg["paths"].items():
            self.cfg["paths"][key] = str(self.base / Path(path).relative_to(BASE))
        self.paths = self.cfg["paths"]
        ensure_libs(self.cfg)
        Path(self.paths["incoming"]).mkdir()
        Path(self.paths["xft_template"]).parent.mkdir()
        shutil.copy2(BASE / "模板" / "新福通通用开票申请单模板.xlsx", self.paths["xft_template"])
        # 仅在自动删除的临时测试目录使用格式测试值，不代表真实税收编码。
        table(self.paths["taxcode_lib"], ["关键词", "税收分类编码", "税率", "单位"],
              [["螺丝", "1" * 19, None, "套"]])

    def input(self, mismatch=False, title=True):
        folder = Path(self.paths["incoming"])
        table(folder / "拼多多订单.xlsx",
              ["订单号", "商品名称", "规格", "数量", "实付金额", "店铺名称"],
              [["ORDER001", "螺丝包", "M30", 1, 10, "测试店"],
               ["ORDER001", "螺丝包", "M40", 2, 20, "测试店"]])
        if title:
            table(folder / "拼多多发票申请.xlsx",
                  ["订单号", "发票抬头", "企业税号", "发票金额", "抬头类型", "店铺名称"],
                  [["ORDER001", "流程测试单位", "TEST_TAX_NO", 31 if mismatch else 30, "企业", "测试店"]])

    def run_flow(self, **kw):
        return run(self.cfg, reporter=SilentReporter(), **kw)

    def report(self, outs):
        return json.loads(Path(outs["处理报告"]).with_suffix(".json").read_text(encoding="utf-8"))

    def test_full_export_register_rerun(self):
        self.input()
        outs = self.run_flow()
        report = self.report(outs)
        self.assertEqual(report["可导出行数"], 2)
        self.assertIn("新福通申请单", outs)
        rows = sheet(outs["新福通申请单"], "导入模板")
        self.assertEqual(rows[2][0], rows[3][0])
        self.assertEqual(rows[2][26], "M30")
        self.assertEqual(rows[3][26], "M40")
        self.assertEqual(rows[2][31] + rows[3][31], 30)
        for row in rows[2:4]:
            self.assertAlmostEqual(row[27] * row[28], row[31], places=2)
            self.assertEqual(row[32], "是")
        again = self.run_flow()
        self.assertNotEqual(outs["新福通申请单"], again["新福通申请单"])
        self.assertNotEqual(rows[2][0], sheet(again["新福通申请单"], "导入模板")[2][0])
        result = self.base / "开票结果.xlsx"
        table(result, ["平台订单号", "发票号码"], [["ORDER001", "INV_TEST_001"]])
        self.assertEqual(mark_issued(self.cfg, result), 1)
        self.assertEqual(mark_issued(self.cfg, result), 0)
        third = self.run_flow()
        self.assertEqual(self.report(third)["明细行数"], 0)
        self.assertNotIn("新福通申请单", third)
        self.assertTrue(list((self.base / "数据" / "备份").glob("*.xlsx")))

    def test_amount_mismatch_blocks_whole_order_and_archive(self):
        self.input(mismatch=True)
        outs = self.run_flow(archive=True)
        self.assertEqual(self.report(outs)["待修正行数"], 2)
        self.assertNotIn("新福通申请单", outs)
        self.assertEqual(len(list(Path(self.paths["incoming"]).glob("*.xlsx"))), 2)

    def test_missing_code_no_import_file(self):
        self.input()
        table(self.paths["taxcode_lib"], ["关键词", "税收分类编码"], [["螺丝", ""]])
        outs = self.run_flow()
        self.assertEqual(self.report(outs)["可导出行数"], 0)
        self.assertNotIn("新福通申请单", outs)

    def test_missing_title_backfill_to_order(self):
        self.input(title=False)
        outs = self.run_flow()
        self.assertIn("待补录", outs)
        wb = openpyxl.load_workbook(outs["待补录"])
        wb.active.cell(4, 10, "补录测试单位")
        wb.active.cell(4, 11, "TEST_TAX_NO")
        wb.save(outs["待补录"])
        wb.close()
        self.assertEqual(import_missing(self.cfg, outs["待补录"]), 1)
        self.assertIn(("平台订单号", "ORDER001"), load_customers(self.paths["customer_lib"]))
        after = self.run_flow()
        self.assertEqual(self.report(after)["可导出行数"], 2)

    def test_register_missing_invoice_atomic(self):
        result = self.base / "result.xlsx"
        table(result, ["平台订单号", "发票号码"], [["A", "INV_A"], ["B", ""]])
        before = Path(self.paths["issued_ledger"]).read_bytes()
        with self.assertRaises(ValueError):
            mark_issued(self.cfg, result)
        self.assertEqual(before, Path(self.paths["issued_ledger"]).read_bytes())
        self.assertEqual(load_issued(self.paths["issued_ledger"]), set())

    def test_register_conflict_atomic(self):
        result = self.base / "result.xlsx"
        table(result, ["平台订单号", "发票号码"], [["A", "INV_A"]])
        mark_issued(self.cfg, result)
        table(result, ["平台订单号", "发票号码"], [["A", "INV_OTHER"]])
        with self.assertRaises(ValueError):
            mark_issued(self.cfg, result)
        self.assertEqual(load_issued(self.paths["issued_ledger"]), {"A"})

    def test_personal_request_clears_old_taxno(self):
        rec = {"平台订单号": "A", "商品名称": "螺丝", "数量": 1, "金额": 10,
               "纳税人识别号": "OLD_COMPANY", "店铺": "测试店"}
        req = {"A": {"title_type": "个人", "title": "张某", "amount": 10}}
        rows, _, _ = build_ledger([rec], self.cfg, {}, [], set(), req)
        self.assertEqual(rows[0]["纳税人识别号"], "")
        self.assertNotIn("缺纳税人识别号", rows[0]["异常标记"])

    def test_net_amount_conversion(self):
        self.cfg["tax"]["amount_is_tax_included"] = False
        rec = {"平台订单号": "A", "商品名称": "螺丝", "数量": 2, "金额": 100, "店铺": "测试店"}
        rows, _, _ = build_ledger([rec], self.cfg, {}, [], set())
        self.assertEqual(rows[0]["价税合计"], 101)
        self.assertEqual(rows[0]["含税单价"], 50.5)

    def test_invoice_types_not_merged(self):
        row = {"发票抬头": "公司", "纳税人识别号": "TEST", "店铺": "店"}
        groups = group_invoices([{**row, "发票类型": "普通"}, {**row, "发票类型": "专用"}], self.cfg["rules"])
        self.assertEqual(len(groups), 2)

    def test_file_failure_prevents_export(self):
        self.input()
        (Path(self.paths["incoming"]) / "bad.xlsx").write_bytes(b"not excel")
        outs = self.run_flow()
        self.assertNotIn("新福通申请单", outs)
        self.assertTrue(self.report(outs)["全局检查"])

    def test_successful_archive(self):
        self.input()
        self.run_flow(archive=True)
        self.assertFalse(list(Path(self.paths["incoming"]).glob("*.xlsx")))
        self.assertEqual(len(list(Path(self.paths["archive"]).rglob("*.xlsx"))), 2)

    def test_conflicting_requests_block(self):
        self.input()
        table(Path(self.paths["incoming"]) / "拼多多另一份发票申请.xlsx",
              ["订单号", "发票抬头", "企业税号", "发票金额", "店铺名称"],
              [["ORDER001", "另一个单位", "OTHER_TAX", 30, "测试店"]])
        outs = self.run_flow()
        self.assertNotIn("新福通申请单", outs)
        self.assertTrue(any("冲突" in x["原因"] for x in self.report(outs)["明细问题"]))

    def test_gui_processes_report(self):
        import tkinter as tk
        import gui
        self.input()
        root = tk.Tk()
        root.withdraw()
        try:
            with patch.object(gui, "load_cfg", side_effect=lambda: copy.deepcopy(self.cfg)), \
                 patch.object(gui, "set_cfg_mode"), patch.object(gui, "BASE", str(self.base)):
                app = gui.App(root)
                app.start()
                deadline = time.monotonic() + 15
                while app.running and time.monotonic() < deadline:
                    root.update()
                    time.sleep(0.01)
                self.assertFalse(app.running, "窗口处理未在期限内完成")
                self.assertIn("处理报告", app.outs)
                self.assertIn("通过 2 行", app.var_pct.get())
        finally:
            root.destroy()
            if 'app' in locals() and app.logger:
                for handler in list(app.logger.handlers):
                    app.logger.removeHandler(handler)
                    handler.close()

    def test_zero_quantity_and_partial_order_block(self):
        row = {"平台订单号": "A", "商品名称": "螺丝", "店铺": "测试店", "数量": 1,
               "价税合计": 10, "税率": .01, "税收分类编码": "1" * 19, "异常标记": ""}
        ready, blocked, _, _ = assess([row.copy(), {**row, "数量": 0}], self.cfg, [], [])
        self.assertEqual(len(ready), 0)
        self.assertEqual(len(blocked), 2)

    def test_erp_number_amount_match(self):
        rec = {"平台订单号": "P", "ERP订单号": "E", "商品名称": "螺丝", "数量": 1, "金额": 10}
        req = {"E": {"title": "测试", "tax_no": "TEST", "amount": 20, "_amount_present": True}}
        rows, _, _ = build_ledger([rec], self.cfg, {}, [], set(), req)
        self.assertIn("申请金额与明细不符", rows[0]["异常标记"])

    def test_not_merge_buyer_keeps_order_lines_together(self):
        row = {"发票抬头": "公司", "纳税人识别号": "TEST", "店铺": "店", "平台订单号": "A"}
        groups = group_invoices([row.copy(), row.copy(), {**row, "平台订单号": "B"}], {"merge_same_buyer": False})
        self.assertEqual(sorted(map(len, groups.values())), [1, 2])


if __name__ == "__main__":
    unittest.main(verbosity=2)
