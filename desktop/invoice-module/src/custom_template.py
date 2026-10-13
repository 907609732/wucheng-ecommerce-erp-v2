# -*- coding: utf-8 -*-
"""
官方模板对齐：如果你把新福通 / 税务局里导出的**官方空白模板**放进「模板」文件夹，
程序会直接按官方表头填数据，避免字段对不上。
"""
import os
import openpyxl
from openpyxl.utils import get_column_letter

from .readers import find_header_row, read_raw

# 表头关键词 -> 我们内部的取值函数
# 顺序很重要：更特殊的放前面（如「税收分类编码」要在「商品编码」之前）
FIELD_RULES = [
    ("tax_code",  ["税收分类编码", "税收编码", "商品和服务税收分类编码", "分类编码"]),
    ("net",       ["不含税金额", "金额（不含税）", "金额"]),
    ("tax",       ["税额"]),
    ("rate",      ["税率"]),
    ("qty",       ["数量"]),
    ("price",     ["单价"]),
    ("unit",      ["计量单位", "单位"]),
    ("spec",      ["规格型号", "规格"]),
    ("name",      ["货物或应税劳务、服务名称", "货物或应税劳务", "项目名称",
                   "服务名称", "商品名称", "货物名称"]),
    ("sku_code",  ["企业商品编码", "商品编码", "货号"]),
    ("title",     ["购买方名称", "购方名称", "抬头", "客户名称", "付款方名称"]),
    ("tax_no",    ["纳税人识别号", "识别号", "税号", "统一社会信用代码"]),
    ("remark",    ["备注"]),
]


def _map_headers(headers):
    """表头 -> {field: 列号}"""
    cmap, used = {}, set()
    for field, kws in FIELD_RULES:
        for kw in kws:
            for i, h in enumerate(headers):
                if i in used or not h:
                    continue
                if kw in str(h):
                    cmap[field] = i
                    used.add(i)
                    break
            if field in cmap:
                break
    return cmap


def _row_values(r, cmap):
    qty = float(r.get("数量", 0) or 0)
    net = float(r.get("不含税金额", 0) or 0)
    return {
        "tax_code": r.get("税收分类编码", ""),
        "net": net,
        "tax": r.get("税额", 0),
        "rate": r.get("税率", 0),
        "qty": qty,
        "price": round(net / qty, 8) if qty else 0,
        "unit": r.get("单位", "件"),
        "spec": r.get("规格型号", ""),
        "name": r.get("开票商品名称") or r.get("商品名称", ""),
        "sku_code": r.get("商品编码", ""),
        "title": r.get("发票抬头", ""),
        "tax_no": r.get("纳税人识别号", ""),
        "remark": r.get("平台订单号", ""),
    }


def render(tmpl_path, groups, cfg, out_path):
    wb = openpyxl.load_workbook(tmpl_path)
    ws = wb.worksheets[0]

    raw = read_raw(tmpl_path)
    hi = find_header_row(raw)
    headers = [str(c).strip() if c else "" for c in raw[hi]]
    cmap = _map_headers(headers)
    if not cmap:
        raise RuntimeError("官方模板里没识别出可填充的列")

    # 表头所在行（Excel 行号从 1 开始）
    excel_header_row = hi + 1
    start_row = excel_header_row + 1

    # 清掉模板里可能存在的示例行
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

    row_i = start_row
    for (title, tax_no, shop, *_), rows in groups.items():
        for r in rows:
            vals = _row_values(r, cmap)
            for field, col in cmap.items():
                ws.cell(row=row_i, column=col + 1, value=vals.get(field, ""))
            row_i += 1

    wb.save(out_path)
    return out_path
