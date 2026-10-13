# -*- coding: utf-8 -*-
"""
流程编排：加载 → 扫描 → 解析匹配 → 汇总 → 导出 → 完成

每一步都通过 reporter 上报状态与日志：
  - 命令行模式打印到屏幕
  - 窗口模式推进队列，由窗口实时显示
  - 两种模式都会写进 日志/运行日志_日期时间.log
"""
import os
import shutil
from datetime import datetime

from .readers import scan_folder
from .requests import load_all
from .ledger import (ensure_libs, load_customers, load_taxcodes, load_issued,
                     build_ledger, group_invoices, find_orphan_requests)
from .reporter import ConsoleReporter, STEP_NAMES
from . import exporters
from .workflow import assess, write_report


def _stamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S_%f")


def run(cfg, archive=False, reporter=None):
    R = reporter or ConsoleReporter()

    base = cfg["_base"]
    paths = cfg["paths"]
    incoming = os.path.join(base, paths["incoming"])
    outdir = os.path.join(base, paths["output"])
    os.makedirs(outdir, exist_ok=True)
    os.makedirs(incoming, exist_ok=True)

    mode = str(cfg.get("mode", "auto")).strip().lower()
    mode_cn = {"erp": "ERP 导出模式", "platform": "平台导出模式"}.get(mode, "自动识别模式")

    # ================= 步骤 1 =================
    R.step(1, STEP_NAMES[0])
    R.detail(f"开票数据模式：{mode_cn}")
    ensure_libs(cfg)
    customers = load_customers(os.path.join(base, paths["customer_lib"]))
    taxcodes = load_taxcodes(os.path.join(base, paths["taxcode_lib"]))
    issued = load_issued(os.path.join(base, paths["issued_ledger"]))
    R.detail(f"抬头库 {len(customers)} 条 | 税收编码库 {len(taxcodes)} 条 | 已开票记录 {len(issued)} 条")
    R.done(1, f"抬头库 {len(customers)} 条 / 编码库 {len(taxcodes)} 条")

    # ================= 步骤 2 =================
    R.step(2, STEP_NAMES[1])
    R.detail(f"扫描文件夹：{incoming}")
    files = scan_folder(incoming)
    if not files:
        R.fail(2, "文件夹里没有找到订单文件（支持 .xlsx/.xls/.csv）")
        R.detail(f"把聚水潭 / 拼多多 / 淘宝导出的表放进「{paths['incoming']}」再运行", "warning")
        R.finish(False, {})
        return None
    R.detail(f"发现 {len(files)} 个待处理文件")
    R.done(2, f"发现 {len(files)} 个文件")

    # ================= 步骤 3 =================
    R.step(3, STEP_NAMES[2])
    details, req_map, fsummary = load_all(files, mode)
    selected_names = {s["file"] for s in fsummary if not s.get("error")}
    selected_files = [f for f in files if os.path.basename(f) in selected_names]
    request_count = len({str(r.get("order_no") or r.get("sale_no") or "") for r in req_map.values()})
    role_cn = {"request": "开票申请", "detail": "订单明细", "both": "申请+明细"}
    for s in fsummary:
        if s.get("error"):
            R.detail(f"× 读取失败 {s['file']}：{s['error']}", "error")
        else:
            R.detail(f"√ {s['file']}　【{s['source']}】{role_cn.get(s['role'], s['role'])}"
                     f"　明细 {s['detail']} 行 / 申请 {s['request']} 条")
    if mode != "auto" and not fsummary:
        R.detail(f"当前是「{mode_cn}」，但投放的文件都不属于该模式，请改 config.yaml 的 mode", "warning")

    if not details and not req_map:
        R.fail(3, "没有解析到有效数据")
        R.finish(False, {})
        return None

    ledger, missing, no_code = build_ledger(details, cfg, customers, taxcodes, issued, req_map)
    orphans = find_orphan_requests(details, req_map)
    ready, blocked, global_issues, problems = assess(ledger, cfg, selected_files, fsummary)
    for issue in global_issues:
        R.detail(issue, "warning")
    R.detail(f"导出前检查：通过 {len(ready)} 行 / 待修正 {len(blocked)} 行")

    R.detail(f"开票申请 {request_count} 个订单 | 订单明细 {len(details)} 行")
    R.detail(f"生成待开票明细 {len(ledger)} 行")
    if missing:
        R.detail(f"! {len(missing)} 行缺发票抬头/税号 → 见「待补录清单」", "warning")
    if no_code:
        R.detail(f"! {no_code} 行缺税收分类编码 → 去「数据/税收编码映射.xlsx」补一次", "warning")
    if orphans:
        R.detail(f"! {len(orphans)} 条开票申请在订单明细里找不到 → 见「漏开票预警」", "warning")

    bad = [r for r in ledger if r.get("金额校验") and r["金额校验"] != "一致"]
    if bad:
        R.detail(f"! {len(bad)} 行「申请金额」与明细合计不一致，已在台账中标红", "warning")

    R.done(3, f"{len(ledger)} 行明细 / {request_count} 个申请订单")

    # ================= 步骤 4 =================
    if not ledger and not missing:
        R.skip(4, "没有需要开票的订单（可能都已开过 / 被规则过滤）")
        outs, stamp = {}, _stamp()
        if orphans:
            R.step(5, STEP_NAMES[4])
            p = os.path.join(outdir, f"漏开票预警_{stamp}.xlsx")
            exporters.write_orphan_requests(orphans, p)
            outs["漏开票预警"] = p
            R.detail(f"√ [漏开票预警] {os.path.basename(p)}")
            R.done(5, f"生成 {len(outs)} 个文件")
        else:
            R.skip(5, "无数据可导出")
        report_path = os.path.join(outdir, f"处理报告_{stamp}.html")
        write_report(report_path, ledger, ready, blocked, global_issues, problems,
                     orphans, fsummary, outs, [])
        outs["处理报告"] = report_path
        R.step(6, STEP_NAMES[5])
        R.done(6, "本次没有需要开票的订单")
        R.finish(True, outs)
        return outs or None

    R.step(4, STEP_NAMES[3])
    groups = group_invoices(ready, cfg["rules"])
    total = sum(float(r.get("价税合计", 0) or 0) for r in ready)
    R.detail(f"按「抬头 + 税号 + 店铺」分组，一个分组 = 一张发票")
    R.done(4, f"{len(groups)} 张发票 / 价税合计 ¥{total:,.2f}")

    # ================= 步骤 5 =================
    R.step(5, STEP_NAMES[4])
    stamp = _stamp()
    outs = {}
    errors = []
    oc = cfg["outputs"]

    def emit(kind, fname, fn):
        p = os.path.join(outdir, f"{fname}_{stamp}.xlsx")
        try:
            fn(p)
            outs[kind] = p
            R.detail(f"√ [{kind}] {os.path.basename(p)}")
        except Exception as e:
            errors.append(f"{kind}生成失败：{e}")
            R.detail(errors[-1], "error")

    if oc.get("ledger", True):
        emit("台账", "待开票台账", lambda p: exporters.write_ledger(ledger, p, no_code=no_code))
    if groups and oc.get("xinfutong_request", True):
        tpl = os.path.join(base, paths.get("xft_template", "模板/新福通通用开票申请单模板.xlsx"))
        try:
            emit("新福通申请单", "新福通开票申请单",
                 lambda p: exporters.write_xinfutong_request(groups, tpl, p, cfg))
        except Exception as e:
            R.detail(f"× 新福通申请单生成失败：{e}", "error")
    if groups and oc.get("zhejiang_etax", True):
        emit("浙江电子税务局", "浙江电子税务局批量开票",
             lambda p: exporters.write_zhejiang_etax(groups, p, cfg))
    if groups and oc.get("xinfutong", True):
        emit("新福通", "新福通清单导入", lambda p: exporters.write_xinfutong(groups, p, cfg))
    if oc.get("missing", True) and missing:
        emit("待补录", "待补录清单", lambda p: exporters.write_missing(missing, p))
    if orphans:
        emit("漏开票预警", "漏开票预警", lambda p: exporters.write_orphan_requests(orphans, p))

    if blocked:
        emit("待修正明细", "待修正明细", lambda p: exporters.write_ledger(blocked, p))
    report_path = os.path.join(outdir, f"处理报告_{stamp}.html")
    write_report(report_path, ledger, ready, blocked, global_issues, problems,
                 orphans, fsummary, outs, errors)
    outs["处理报告"] = report_path

    R.done(5, f"生成 {len(outs)} 个文件")

    # ================= 步骤 6 =================
    R.step(6, STEP_NAMES[5])
    if archive and not (blocked or orphans or errors or global_issues):
        adir = os.path.join(base, paths["archive"], stamp)
        os.makedirs(adir, exist_ok=True)
        for f in selected_files:
            shutil.move(f, os.path.join(adir, os.path.basename(f)))
        R.detail(f"已归档 {len(selected_files)} 个本模式原始文件 → {adir}")
    else:
        R.detail("原始文件保留在「订单投放」；实际开票后须登记发票号码，才能在下次处理时跳过订单")
    R.done(6, "全部完成")
    R.finish(not errors, outs)

    return outs


def import_missing(cfg, source_path=None):
    """把「待补录清单」里手工填好的抬头/税号回写进抬头库，以后自动匹配"""
    import glob
    import openpyxl
    base = cfg["_base"]
    outdir = os.path.join(base, cfg["paths"]["output"])
    lib_path = os.path.join(base, cfg["paths"]["customer_lib"])

    files = ([source_path] if source_path else
             sorted(glob.glob(os.path.join(outdir, "待补录清单_*.xlsx")), reverse=True))
    if not files:
        print("没有找到待补录清单，请先运行一次生成。")
        return 0

    wb = openpyxl.load_workbook(files[0], data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 4:
        print("清单为空。")
        return 0

    lib = openpyxl.load_workbook(lib_path)
    lw = lib.worksheets[0]
    exist = set()
    for r in lw.iter_rows(min_row=2, values_only=True):
        if r and len(r) > 1 and r[1]:
            exist.add((str(r[0] or "").strip(), str(r[1]).strip()))

    added = 0
    for r in rows[3:]:
        if not r or len(r) < 11:
            continue
        buyer, phone, receiver = str(r[0] or "").strip(), str(r[1] or "").strip(), str(r[2] or "").strip()
        title, taxno = str(r[9] or "").strip(), str(r[10] or "").strip()
        if not title or not taxno:
            continue
        addr = str(r[11] or "").strip() if len(r) > 11 else ""
        email = str(r[12] or "").strip() if len(r) > 12 else ""
        order_no = str(r[3] or "").strip()
        # 始终按订单号回写，防止同一买家为不同单位申请时错套抬头。
        key = order_no
        ktype = "平台订单号"
        if not key or (ktype, key) in exist:
            continue
        lw.append([ktype, key, title, taxno, addr, "", "", email, "待补录回写"])
        exist.add((ktype, key))
        added += 1

    if added:
        _save_library(lib, lib_path)
    lib.close()
    print(f"已回写 {added} 条抬头到抬头库：{lib_path}")
    return added


def mark_issued(cfg, path):
    """把已开完的发票结果（含发票号码）并入已开票记录，用于去重"""
    import openpyxl
    base = cfg["_base"]
    dst = os.path.join(base, cfg["paths"]["issued_ledger"])
    from .readers import read_raw
    rows = read_raw(path)
    if not rows:
        return 0
    hi = next((i for i, row in enumerate(rows[:15]) if "平台订单号" in row and "发票号码" in row), None)
    if hi is None:
        raise ValueError("结果表必须有「平台订单号」「发票号码」两列；仅登记实际开票成功的订单")
    head = rows[hi]
    i_no = next((i for i, h in enumerate(head) if "平台订单号" in h), None)
    i_inv = next((i for i, h in enumerate(head) if "发票号码" in h), None)
    ensure_libs(cfg)

    lib = openpyxl.load_workbook(dst)
    lw = lib.worksheets[0]
    existing = {str(r[0] or "").strip(): str(r[2] or "").strip()
                for r in lw.iter_rows(min_row=2, values_only=True) if r and r[0]}
    additions = {}
    n = 0
    for r in rows[hi + 1:]:
        v = str(r[i_no] or "").strip() if i_no < len(r) else ""
        if not v:
            continue
        inv = str(r[i_inv] or "").strip() if (i_inv is not None and i_inv < len(r)) else ""
        if not inv:
            raise ValueError(f"订单 {v} 缺发票号码：整份文件未登记，请补全实际开票结果")
        import re
        for order in re.split(r"[,，;；\s]+", v):
            old = existing.get(order) or additions.get(order)
            if old and old != inv:
                raise ValueError(f"订单 {order} 已登记不同的发票号码，请人工核对")
            if not existing.get(order) and order not in additions:
                additions[order] = inv
    for order, inv in additions.items():
        lw.append([order, "", inv, datetime.now().strftime("%Y-%m-%d"), "", ""])
        n += 1
    if n:
        _save_library(lib, dst)
    lib.close()
    print(f"已登记 {n} 条已开票记录（之后再跑就不会重复开票了）")
    return n


def _save_library(wb, path):
    """保存前备份；写入失败时原库保持完整。"""
    backup_dir = os.path.join(os.path.dirname(path), "备份")
    os.makedirs(backup_dir, exist_ok=True)
    if os.path.exists(path):
        shutil.copy2(path, os.path.join(backup_dir, f"{_stamp()}_{os.path.basename(path)}"))
    tmp = path + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)
