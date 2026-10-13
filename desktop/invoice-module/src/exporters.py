# -*- coding: utf-8 -*-
"""
导出器：生成
  1) 待开票台账（主表，字段最全，人看的）
  2) 浙江省电子税务局《批量导入开具发票模板》
  3) 新福通 清单导入模板
  4) 待补录清单（缺抬头/税号/编码）
"""
import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10.5)
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(size=9, color="7F7F7F")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _w(ws, cols_width):
    for i, w in enumerate(cols_width, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header(ws, cols, width_map=None, freeze="A2"):
    ws.append(cols)
    for c in ws[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    if width_map:
        _w(ws, width_map)
    else:
        _w(ws, [max(10, min(30, len(str(c)) * 2.2 + 4)) for c in cols])
    ws.freeze_panes = freeze
    ws.auto_filter.ref = None


def _autofilter(ws, ncols, nrows):
    if nrows > 0:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"


def _fmt_rows(ws, money_cols=(), start=2):
    for row in ws.iter_rows(min_row=start):
        for c in row:
            c.border = BORDER
            if c.column in money_cols and isinstance(c.value, (int, float)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")


# ===============================================================
# 1. 待开票台账
# ===============================================================
LEDGER_COLS = ["来源", "店铺", "平台订单号", "ERP订单号", "子订单号", "买家账号", "收货人",
               "手机号", "支付时间", "商品编码", "商品名称", "开票商品名称", "规格型号",
               "单位", "数量", "含税单价", "价税合计", "不含税金额", "税率", "税额",
               "税收分类编码", "发票抬头", "纳税人识别号", "地址电话", "开户行",
               "银行账号", "交付邮箱", "快递公司", "快递单号", "买家留言", "卖家备注",
               "匹配方式", "抬头类型", "发票类型", "开票申请来源", "申请金额", "金额校验",
               "异常标记", "来源文件"]
LEDGER_WIDTH = [10, 14, 20, 14, 16, 16, 10, 14, 18, 14, 26, 22, 14, 6, 6, 10, 11,
                11, 7, 9, 14, 24, 20, 20, 16, 18, 18, 12, 18, 18, 18, 10, 10, 18,
                14, 11, 24, 18, 18]


def write_ledger(rows, path, no_code=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "待开票台账"
    _header(ws, LEDGER_COLS, LEDGER_WIDTH)

    for r in rows:
        ws.append([r.get(c, "") for c in LEDGER_COLS])
        # 税率显示成 13%
        cell = ws.cell(row=ws.max_row, column=LEDGER_COLS.index("税率") + 1)
        try:
            cell.value = f"{float(r.get('税率', 0)) * 100:g}%"
        except (TypeError, ValueError):
            pass

    n = len(rows)
    money_cols = {LEDGER_COLS.index(c) + 1 for c in
                  ("含税单价", "价税合计", "不含税金额", "税额", "申请金额")}
    _fmt_rows(ws, money_cols=money_cols)
    _autofilter(ws, len(LEDGER_COLS), n)

    # 异常行标黄；申请金额与明细对不上的标红加粗
    if n:
        warn_col = LEDGER_COLS.index("异常标记") + 1
        chk_col = LEDGER_COLS.index("金额校验") + 1
        amt_col = LEDGER_COLS.index("价税合计") + 1
        for i in range(2, n + 2):
            if ws.cell(row=i, column=warn_col).value:
                for j in range(1, len(LEDGER_COLS) + 1):
                    ws.cell(row=i, column=j).fill = WARN_FILL
            chk = str(ws.cell(row=i, column=chk_col).value or "")
            if chk and chk != "一致":
                ws.cell(row=i, column=chk_col).font = Font(bold=True, color="C00000")
                ws.cell(row=i, column=amt_col).font = Font(bold=True, color="C00000")

    # 汇总
    ws2 = wb.create_sheet("汇总")
    total = sum(float(r.get("价税合计", 0) or 0) for r in rows)
    net = sum(float(r.get("不含税金额", 0) or 0) for r in rows)
    tax = sum(float(r.get("税额", 0) or 0) for r in rows)
    ws2["A1"] = "待开票汇总"
    ws2["A1"].font = TITLE_FONT
    info = [
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("明细行数", n),
        ("发票张数", len({(str(r.get("发票抬头", "")), str(r.get("纳税人识别号", "")), str(r.get("店铺", ""))) for r in rows})),
        ("价税合计", round(total, 2)),
        ("不含税金额合计", round(net, 2)),
        ("税额合计", round(tax, 2)),
        ("缺税收分类编码行数", no_code),
    ]
    if no_code:
        ws2.cell(row=len(info) + 3, column=1).fill = WARN_FILL
        ws2.cell(row=len(info) + 3, column=2).fill = WARN_FILL
        ws2.cell(row=len(info) + 4, column=1,
                 value="→ 打开「数据/税收编码映射.xlsx」，按商品关键词补税收分类编码，补一次永久生效").font = Font(size=9, color="C00000")
    for i, (k, v) in enumerate(info, start=3):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws2.cell(row=i, column=2, value=v)
        if isinstance(v, float):
            c.number_format = "#,##0.00"
    _w(ws2, [18, 24])

    wb.save(path)
    return path


# ===============================================================
# 2. 浙江省电子税务局《批量导入开具发票模板》
# ===============================================================
ETAX_BASE_COLS = [
    "序号", "发票类型", "购买方名称", "购买方识别号（统一社会信用代码）",
    "购买方地址", "购买方电话", "购买方开户银行", "购买方银行账号",
    "购买方邮箱", "购买方手机号", "备注", "价税合计",
]
ETAX_BASE_WIDTH = [6, 14, 26, 24, 24, 14, 18, 20, 22, 14, 22, 12]

ETAX_DETAIL_COLS = [
    "序号", "项目名称", "商品和服务税收分类编码", "规格型号", "单位",
    "数量", "单价（不含税）", "金额（不含税）", "税率", "税额",
    "折扣金额", "折扣税额", "含税标志",
]
ETAX_DETAIL_WIDTH = [6, 28, 22, 14, 8, 8, 12, 12, 8, 10, 10, 10, 10]


def write_zhejiang_etax(groups, path, cfg):
    wb = openpyxl.Workbook()

    # ---- Sheet1 发票基本信息 ----
    ws = wb.active
    ws.title = "1-发票基本信息"
    ws["A1"] = "填表说明：本页一张发票一行；「序号」必须与「2-发票明细信息」中的序号一致，用于关联明细。"
    ws["A1"].font = NOTE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ETAX_BASE_COLS))
    ws.append([])
    ws.append(ETAX_BASE_COLS)
    hdr_row = 3
    for c in ws[hdr_row]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _w(ws, ETAX_BASE_WIDTH)

    # ---- Sheet2 发票明细信息 ----
    ws2 = wb.create_sheet("2-发票明细信息")
    ws2["A1"] = ("填表说明：本页一行一个商品明细；「序号」填对应发票在基本信息页中的序号。"
                 "金额与单价均为【不含税】，含税标志固定填 0；发票类型可按实际改成电子专用发票。")
    ws2["A1"].font = NOTE_FONT
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ETAX_DETAIL_COLS))
    ws2.append([])
    ws2.append(ETAX_DETAIL_COLS)
    for c in ws2[3]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _w(ws2, ETAX_DETAIL_WIDTH)

    # 明细页填的是「不含税」金额/单价，所以含税标志固定为 0
    flag = 0
    default_email = cfg.get("delivery", {}).get("default_email", "")
    default_phone = cfg.get("delivery", {}).get("default_phone", "")

    idx = 0
    for (title, tax_no, shop, *_), rows in groups.items():
        idx += 1
        total = round(sum(float(r.get("价税合计", 0) or 0) for r in rows), 2)
        addr = next((str(r.get("地址电话", "")) for r in rows if r.get("地址电话")), "")
        email = next((str(r.get("交付邮箱", "")) for r in rows if r.get("交付邮箱")), default_email)
        phone = next((str(r.get("手机号", "")) for r in rows if r.get("手机号")), default_phone)
        bank = next((str(r.get("开户行", "")) for r in rows if r.get("开户行")), "")
        bacc = next((str(r.get("银行账号", "")) for r in rows if r.get("银行账号")), "")
        notes = [str(r.get("平台订单号", "")) for r in rows if r.get("平台订单号")]
        remark = "订单号：" + "，".join(dict.fromkeys(notes))[:80]

        ws.append([idx, "电子发票（普通发票）", title, tax_no, addr, "", bank, bacc,
                   email, phone, remark, total])

        for r in rows:
            qty = float(r.get("数量", 0) or 0)
            net = float(r.get("不含税金额", 0) or 0)
            rate = float(r.get("税率", 0) or 0)
            unit_price = round(net / qty, 8) if qty else 0
            ws2.append([idx, r.get("开票商品名称") or r.get("商品名称"),
                        r.get("税收分类编码", ""), r.get("规格型号", ""),
                        r.get("单位", "件"), qty, unit_price, net, rate,
                        r.get("税额", 0), 0, 0, flag])

    _fmt_rows(ws, money_cols={12}, start=hdr_row + 1)
    _fmt_rows(ws2, money_cols={7, 8, 10, 11, 12}, start=4)
    # 税率格式
    for row in ws2.iter_rows(min_row=4, min_col=9, max_col=9):
        for c in row:
            c.number_format = "0.00"

    # ---- Sheet3/4 ----
    ws3 = wb.create_sheet("3-特定业务信息")
    ws3["A1"] = "本店为普通货物销售，一般无需填写。若涉及建筑服务、货物运输、不动产销售等特定业务，请按税务局模板说明填写。"
    ws3["A1"].font = NOTE_FONT
    ws3.column_dimensions["A"].width = 90

    ws4 = wb.create_sheet("4-附加要素信息")
    ws4["A1"] = "无需填写时留空即可。"
    ws4["A1"].font = NOTE_FONT
    ws4.column_dimensions["A"].width = 60

    wb.save(path)
    return path


# ===============================================================
# 3. 新福通 清单导入模板
# ===============================================================
XFT_HEAD_COLS = ["序号", "购买方名称", "纳税人识别号", "地址、电话", "开户行及账号",
                 "备注", "价税合计", "不含税金额", "税额"]
XFT_HEAD_WIDTH = [6, 28, 22, 28, 28, 30, 12, 12, 12]

# 航信/百望系开票软件通用清单导入字段
XFT_DETAIL_COLS = ["序号", "货物或应税劳务、服务名称", "计量单位", "规格型号", "数量",
                   "金额", "税率", "商品税目", "折扣金额", "税额", "折扣税额", "折扣率",
                   "单价", "价格方式", "税收分类编码", "企业商品编码",
                   "使用优惠政策标识", "零税率标识", "优惠政策说明", "中外合作油气田标识"]
XFT_DETAIL_WIDTH = [6, 30, 10, 14, 8, 12, 8, 12, 10, 12, 10, 8, 10, 10, 18, 16, 14, 12, 16, 16]


def write_xinfutong(groups, path, cfg):
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "发票抬头信息"
    ws["A1"] = ("说明：本页是每张发票的购买方信息，供你按抬头在新福通中逐张填开；"
                "「序号」对应「清单明细」页，把同一序号的明细行导入到该张发票的清单里。")
    ws["A1"].font = NOTE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(XFT_HEAD_COLS))
    ws.append([])
    ws.append(XFT_HEAD_COLS)
    for c in ws[3]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _w(ws, XFT_HEAD_WIDTH)

    ws2 = wb.create_sheet("清单明细")
    ws2["A1"] = ("说明：本页为可直接导入开票软件「清单」的标准字段。"
                 "金额/单价均为【不含税】，价格方式固定 0；优惠政策类字段默认 0。")
    ws2["A1"].font = NOTE_FONT
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(XFT_DETAIL_COLS))
    ws2.append([])
    ws2.append(XFT_DETAIL_COLS)
    for c in ws2[3]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _w(ws2, XFT_DETAIL_WIDTH)

    # 金额/单价均按不含税填写，故价格方式固定 0（0=不含税，1=含税）
    price_mode = 0

    idx = 0
    for (title, tax_no, shop, *_), rows in groups.items():
        idx += 1
        total = round(sum(float(r.get("价税合计", 0) or 0) for r in rows), 2)
        net = round(sum(float(r.get("不含税金额", 0) or 0) for r in rows), 2)
        tax = round(sum(float(r.get("税额", 0) or 0) for r in rows), 2)
        addr = next((str(r.get("地址电话", "")) for r in rows if r.get("地址电话")), "")
        bank = next((str(r.get("开户行", "")) for r in rows if r.get("开户行")), "")
        bacc = next((str(r.get("银行账号", "")) for r in rows if r.get("银行账号")), "")
        bankstr = (bank + " " + bacc).strip()
        notes = [str(r.get("平台订单号", "")) for r in rows if r.get("平台订单号")]
        remark = "订单号：" + "，".join(dict.fromkeys(notes))[:80]

        ws.append([idx, title, tax_no, addr, bankstr, remark, total, net, tax])

        for r in rows:
            qty = float(r.get("数量", 0) or 0)
            net_amt = float(r.get("不含税金额", 0) or 0)
            rate = float(r.get("税率", 0) or 0)
            unit_price = round(net_amt / qty, 8) if qty else 0
            ws2.append([
                idx,
                r.get("开票商品名称") or r.get("商品名称"),
                r.get("单位", "件"),
                r.get("规格型号", ""),
                qty,
                net_amt,          # 金额（不含税）
                rate,             # 税率
                "",               # 商品税目（非必填）
                0,                # 折扣金额
                r.get("税额", 0),  # 税额
                0,                # 折扣税额
                0,                # 折扣率
                unit_price,       # 单价（不含税）
                price_mode,
                r.get("税收分类编码", ""),
                r.get("商品编码", ""),
                0, 0, "", 0,
            ])

    _fmt_rows(ws, money_cols={7, 8, 9}, start=4)
    _fmt_rows(ws2, money_cols={6, 10, 13}, start=4)
    for row in ws2.iter_rows(min_row=4, min_col=7, max_col=7):
        for c in row:
            c.number_format = "0.00"

    ws3 = wb.create_sheet("使用说明")
    steps = [
        "1. 打开新福通开票软件 → 发票管理 → 发票填开管理 → 增值税普通发票（或专用发票）",
        "2. 在填开界面点击「清单」→「导入」→ 选择本文件的「清单明细」页（需先另存为该页单独的 xls/xlsx）",
        "3. 导入前请先在开票软件里「系统设置 → 商品编码」中导入一次税收分类编码，否则编码为空会导入失败",
        "4. 「税收分类编码」列为必填。请在开票软件中查询本店商品对应的编码后，回填到「数据/税收编码映射.xlsx」",
        "5. 若新福通版本要求的是官方模板，请在开票软件里导出一次官方模板，把它放到本项目的「模板」文件夹，",
        "   程序下次运行会自动对齐官方表头。",
    ]
    for i, s in enumerate(steps, start=1):
        ws3.cell(row=i, column=1, value=s)
    ws3.column_dimensions["A"].width = 100

    wb.save(path)
    return path


# ===============================================================
# 4. 待补录清单
# ===============================================================
MISSING_COLS = ["买家账号", "手机号", "收货人", "平台订单号", "店铺", "商品名称",
                "价税合计", "现有抬头", "待补录原因", "→请填抬头", "→请填税号",
                "→请填地址电话", "→请填邮箱"]
MISSING_WIDTH = [16, 14, 12, 20, 14, 26, 11, 20, 18, 22, 22, 20, 20]


def write_missing(rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "待补录"
    ws["A1"] = ("以下订单缺抬头/税号。在黄色列补填后，点击窗口的「补录回写」选择此文件，"
                "再重新处理。税收编码请在商品编码库补充。")
    ws["A1"].font = Font(size=10, color="C00000")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(MISSING_COLS))
    ws.append([])
    ws.append(MISSING_COLS)
    for c in ws[3]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _w(ws, MISSING_WIDTH)

    seen = set()
    for r in rows:
        key = (str(r.get("买家账号", "")), str(r.get("平台订单号", "")))
        if key in seen:
            continue
        seen.add(key)
        ws.append([r.get("买家账号", ""), r.get("手机号", ""), r.get("收货人", ""),
                   r.get("平台订单号", ""), r.get("店铺", ""), r.get("商品名称", ""),
                   r.get("价税合计", 0), r.get("发票抬头", ""), r.get("待补录原因", ""),
                   "", "", "", ""])

    for row in ws.iter_rows(min_row=4, min_col=10, max_col=13):
        for c in row:
            c.fill = WARN_FILL
    _fmt_rows(ws, money_cols={7}, start=4)
    wb.save(path)
    return path


# ===============================================================
# 4.5 新福通「开票申请单」导入模板（贴官方模板，一键导入统一开票）
# ===============================================================
# 官方模板 36 列，第 1 行是导入说明，第 2 行是表头，第 3 行起是数据
XFT_REQ_SHEET = "导入模板"
XFT_REQ_DATA_START = 3
XFT_REQ_NCOL = 36

# 列号(1-based) -> 用途，方便对照排查
XFT_REQ_MAP = {
    "bill_no": 1, "customer": 2, "buyer_name": 3, "buyer_taxno": 4,
    "buyer_addr": 5, "buyer_phone": 6, "buyer_bank": 7, "buyer_acct": 8,
    "is_natural": 9, "seller_legal": 10,
    "seller_addr": 11, "seller_phone": 12, "seller_bank": 13, "seller_acct": 14,
    "inv_type": 15, "special_biz": 16, "reduce_tax": 17, "remark": 18,
    "show_b_acct": 19, "show_b_addr": 20, "show_s_acct": 21, "show_s_addr": 22,
    "email": 23, "tax_code": 24, "inv_name": 25, "unit": 26, "spec": 27,
    "qty": 28, "price": 29, "rate": 30, "policy": 31, "amount": 32,
    "tax_incl": 33, "disc_rate": 34, "disc_amt": 35, "goods_name": 36,
}


def _rate_text(rate):
    """0.01 -> '1%'；模板要求带百分号的文本"""
    try:
        r = float(rate or 0)
    except (TypeError, ValueError):
        r = 0.0
    return f"{r * 100:g}%"


def write_xinfutong_request(groups, template_path, out_path, cfg):
    """
    把数据贴进新福通官方「通用开票申请单」模板。

    官方规则：**开票单单据号相同的数据，会开在同一张发票上** —— 所以
    一个买家分组 = 一个单据号 = 一张发票。

    金额口径：模板要求「金额 = 数量 × 单价」，所以这里统一按【含税】填，
    并把「是否含税」写成「是」，避免不含税价四舍五入后等式对不上导致导入失败。
    """
    import shutil

    if not os.path.exists(template_path):
        raise FileNotFoundError(
            f"找不到新福通模板：{template_path}\n"
            f"请把官方模板放到「模板」文件夹，文件名：新福通通用开票申请单模板.xlsx")

    shutil.copyfile(template_path, out_path)
    wb = openpyxl.load_workbook(out_path)
    if XFT_REQ_SHEET not in wb.sheetnames:
        raise RuntimeError(f"模板里没找到「{XFT_REQ_SHEET}」工作表")
    ws = wb[XFT_REQ_SHEET]
    last_row = max(ws.max_row, XFT_REQ_DATA_START)

    seller_name = str(cfg.get("seller", {}).get("name", "")).strip()
    prefix = datetime.now().strftime("%Y%m%d%H%M%S%f")

    r = XFT_REQ_DATA_START
    seq = 0
    for (title, tax_no, shop, *_), rows in groups.items():
        seq += 1
        bill_no = f"{prefix}{seq:04d}"

        first = rows[0]
        addr = next((str(x.get("地址电话", "")) for x in rows if x.get("地址电话")), "")
        email = next((str(x.get("交付邮箱", "")) for x in rows if x.get("交付邮箱")), "")
        title_type = str(first.get("抬头类型", ""))
        inv_raw = str(first.get("发票类型", ""))
        inv_type = ("电子发票（增值税专用发票）" if "专用" in inv_raw
                    else "电子发票（普通发票）")
        is_natural = "是" if "个人" in title_type else "否"

        notes = [str(x.get("平台订单号", "")) for x in rows if x.get("平台订单号")]
        remark = "，".join(dict.fromkeys(notes))[:190]

        for x in rows:
            qty = float(x.get("数量", 0) or 0)
            amount = round(float(x.get("价税合计", 0) or 0), 2)  # 含税金额
            unit_price = round(amount / qty, 8) if qty else 0

            vals = {
                "bill_no": bill_no,
                "customer": "",                       # 留空，直接填购买方信息
                "buyer_name": title,
                "buyer_taxno": tax_no,
                "buyer_addr": addr,
                "buyer_phone": "",
                "buyer_bank": "",
                "buyer_acct": "",
                "is_natural": is_natural,
                "seller_legal": seller_name,          # 必填
                "seller_addr": "",
                "seller_phone": "",
                "seller_bank": "",
                "seller_acct": "",
                "inv_type": inv_type,                 # 必填
                "special_biz": "",
                "reduce_tax": "",
                "remark": remark,
                "show_b_acct": "", "show_b_addr": "", "show_s_acct": "", "show_s_addr": "",
                "email": email,
                "tax_code": x.get("税收分类编码", ""),
                "inv_name": x.get("开票商品名称") or x.get("商品名称", ""),
                "unit": x.get("单位", ""),
                "spec": x.get("规格型号", ""),
                "qty": qty,
                "price": unit_price,                  # 含税单价
                "rate": _rate_text(x.get("税率", 0)),
                "policy": "不使用优惠政策",             # 取下拉列表里的确切值，避免留空被判无效
                "amount": amount,                     # 含税金额
                "tax_incl": "是",
                "disc_rate": "",
                "disc_amt": "",
                "goods_name": x.get("商品名称", ""),
            }
            for key, col in XFT_REQ_MAP.items():
                ws.cell(row=r, column=col, value=vals.get(key, ""))
            r += 1

    # 清掉模板自带的示例数据（保留格式与数据验证）
    for rr in range(r, last_row + 1):
        for c in range(1, XFT_REQ_NCOL + 1):
            ws.cell(row=rr, column=c).value = None

    wb.save(out_path)
    return out_path


# ===============================================================
# 5. 漏开票预警：有买家开了申请，但订单明细里找不到
# ===============================================================
ORPHAN_COLS = ["申请来源", "平台", "店铺", "线上订单号", "销售单号", "抬头类型",
               "发票抬头", "纳税人识别号", "申请金额", "交付邮箱", "申请时间",
               "备注", "处理建议"]
ORPHAN_WIDTH = [12, 10, 22, 22, 12, 10, 26, 22, 11, 20, 19, 18, 30]


def write_orphan_requests(rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "有申请无订单"
    ws["A1"] = ("【重要】以下买家已经提交了开票申请，但在你导入的订单明细里找不到对应订单。"
                "原因通常是：订单明细导出时选的日期范围没覆盖到，或该店铺的销售单还没导出。"
                "请补导订单明细，否则这些票会被漏掉。")
    ws["A1"].font = Font(bold=True, size=10, color="C00000")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ORPHAN_COLS))
    ws.append([])
    ws.append(ORPHAN_COLS)
    for c in ws[3]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _w(ws, ORPHAN_WIDTH)

    for r in rows:
        ws.append([
            r.get("req_source", ""), r.get("platform", ""), r.get("shop", ""),
            r.get("order_no", ""), r.get("sale_no", ""), r.get("title_type", ""),
            r.get("title", ""), r.get("tax_no", ""), r.get("amount", 0),
            r.get("email", ""), r.get("apply_time", ""), r.get("remark", ""),
            "补导该订单的销售单/订单明细",
        ])
    _fmt_rows(ws, money_cols={9}, start=4)
    wb.save(path)
    return path
